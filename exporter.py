"""把分析报告与历史会话导出为 Word（.docx）文档。"""
import datetime as dt
import io
import re

from docx import Document
from docx.shared import Pt

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_TABLE_ROW_RE = re.compile(r"^\|(.+)\|\s*$")


def _add_rich_text(paragraph, text: str) -> None:
    """按 **加粗** 拆分文本为多个 run。"""
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        run = paragraph.add_run(m.group(1))
        run.bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _is_table_separator(cells: list) -> bool:
    return all(re.fullmatch(r":?-{2,}:?", c.strip() or "") for c in cells)


def _split_row(line: str) -> list:
    return [c.strip() for c in _TABLE_ROW_RE.match(line).group(1).split("|")]


def render_markdown(doc: Document, text: str) -> None:
    """把 Markdown 文本渲染进文档：支持标题、列表、表格、加粗，其余按段落处理。"""
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
            i += 1
            continue
        if _TABLE_ROW_RE.match(line):
            rows = []
            while i < len(lines) and _TABLE_ROW_RE.match(lines[i].rstrip()):
                rows.append(_split_row(lines[i].rstrip()))
                i += 1
            rows = [r for r in rows if not _is_table_separator(r)]
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c in range(cols):
                        cell = table.cell(r, c)
                        cell.text = ""
                        _add_rich_text(cell.paragraphs[0], row[c] if c < len(row) else "")
                        if r == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
            continue
        m = re.match(r"^[-*]\s+(.*)$", line)
        if m:
            _add_rich_text(doc.add_paragraph(style="List Bullet"), m.group(1))
            i += 1
            continue
        m = re.match(r"^\d+[.、]\s*(.*)$", line)
        if m:
            _add_rich_text(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue
        _add_rich_text(doc.add_paragraph(), line)
        i += 1


def _fmt_time(iso: str) -> str:
    return str(iso or "").replace("T", " ")


def build_report_docx(target: str, latest_report: str, sessions: list) -> bytes:
    """生成「最新报告 + 历史会话」的 Word 文档，返回字节内容。"""
    doc = Document()
    doc.styles["Normal"].font.name = "Microsoft YaHei"
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading(f"邮件分析报告：{target}", level=0)
    doc.add_paragraph(f"导出时间：{dt.datetime.now():%Y-%m-%d %H:%M}")

    doc.add_heading("最新报告", level=1)
    if latest_report:
        render_markdown(doc, latest_report)
    else:
        doc.add_paragraph("（暂无报告）")

    if sessions:
        doc.add_heading("历史会话", level=1)
        for i, sess in enumerate(sessions, 1):
            doc.add_heading(f"第 {i} 次分析 · {_fmt_time(sess.get('started_at'))}", level=2)
            for m in sess.get("messages") or []:
                label = "用户" if m.get("role") == "user" else "助手"
                p = doc.add_paragraph()
                run = p.add_run(f"【{label} · {_fmt_time(m.get('time'))}】")
                run.bold = True
                render_markdown(doc, m.get("content", ""))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
